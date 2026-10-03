"""DOCX renderer (python-docx) — editable Word copies of every printable."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from . import icons
from .markup import tokenize
from .model import (
    Block, Box, Bullets, Doc, Group, Heading, Lines, PageBreak, Para, Rule, Spacer, Table,
)

FONT = "Arial"
PRIMARY = RGBColor(0x1B, 0x6C, 0xA8)
ACCENT = RGBColor(0xE8, 0x87, 0x1E)
KEYRED = RGBColor(0xB3, 0x26, 0x1E)

BOX_FILL = {"note": "EAF3FB", "tip": "EAF7EA", "bank": "FFF6E0", "grammar": "F2EBFA",
            "answer": "FDECEA", "script": "FDECEA", "draw": "FFFFFF"}
BOX_EDGE = {"note": "1B6CA8", "tip": "2E7D32", "bank": "E8871E", "grammar": "6A3FA0",
            "answer": "B3261E", "script": "B3261E", "draw": "9FB3C8"}


TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]


def _local(el) -> str:
    return el.tag.rsplit("}", 1)[-1]


def _insert_ordered(parent, child, order: list[str]) -> None:
    """Insert ``child`` into ``parent`` respecting the OOXML schema sequence (Word is strict)."""
    name = _local(child)
    for existing in list(parent):
        if _local(existing) == name:
            parent.remove(existing)
    rank = order.index(name)
    for existing in parent:
        if _local(existing) in order and order.index(_local(existing)) > rank:
            existing.addprevious(child)
            return
    parent.append(child)


def _shade(cell, hex_fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    _insert_ordered(cell._tc.get_or_add_tcPr(), shd, TCPR_ORDER)


def _borders(cell, color: str = "9FB3C8", sz: int = 6) -> None:
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    _insert_ordered(cell._tc.get_or_add_tcPr(), borders, TCPR_ORDER)


def _bottom_border(paragraph, color: str = "9FB3C8") -> None:
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    _insert_ordered(paragraph._p.get_or_add_pPr(), pbdr, PPR_ORDER)


def _runs(paragraph, text: str, size: float = 11, bold: bool = False, color: RGBColor | None = None,
          italic: bool = False) -> None:
    for tok in tokenize(text):
        if tok.kind == "br":
            paragraph.add_run().add_break()
            continue
        if tok.kind == "icon":
            path = icons.png(tok.value)
            if path is not None:
                paragraph.add_run().add_picture(str(path), width=Pt(tok.size or size + 3))
            continue
        run = paragraph.add_run(tok.value)
        if tok.kind == "ans":
            color = KEYRED
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(size)
        run.bold = bold or tok.kind in ("bold", "ans")
        run.italic = italic
        if color is not None:
            run.font.color.rgb = color


def _first_par(container):
    """Re-use the empty first paragraph of a fresh table cell."""
    if hasattr(container, "paragraphs") and container.paragraphs and not container.paragraphs[0].text \
            and getattr(container, "_used", False) is False and container.__class__.__name__ == "_Cell":
        container._used = True  # type: ignore[attr-defined]
        return container.paragraphs[0]
    return container.add_paragraph()


def _emit(container, blocks: list[Block], doc_width_mm: float) -> None:
    for b in blocks:
        _block(container, b, doc_width_mm)


def _block(container, b: Block, width_mm: float) -> None:
    if isinstance(b, Heading):
        p = _first_par(container)
        sizes = {1: 18, 2: 14, 3: 12}
        _runs(p, b.text, size=sizes.get(b.level, 12), bold=True,
              color=PRIMARY if b.level < 3 else ACCENT)
        p.paragraph_format.space_before = Pt(10 if b.level > 1 else 6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
    elif isinstance(b, Para):
        p = _first_par(container)
        size = {"small": 9, "note": 10, "instruction": 12}.get(b.style, 11)
        _runs(p, b.text, size=size, bold=b.style == "instruction", italic=b.style == "note")
        if b.style == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif b.style == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(4)
        if b.style == "instruction":
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
    elif isinstance(b, Bullets):
        for it in b.items:
            p = container.add_paragraph(style="List Number" if b.ordered else "List Bullet")
            _runs(p, it, size=11)
            p.paragraph_format.space_after = Pt(2)
    elif isinstance(b, Spacer):
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(b.h * 2.2)
    elif isinstance(b, PageBreak):
        container.add_page_break() if hasattr(container, "add_page_break") else None
    elif isinstance(b, Rule):
        p = container.add_paragraph()
        _bottom_border(p)
    elif isinstance(b, Lines):
        for _ in range(b.n):
            p = container.add_paragraph()
            p.paragraph_format.space_before = Pt(b.height * 2.3)
            p.paragraph_format.space_after = Pt(0)
            _bottom_border(p)
    elif isinstance(b, Group):
        _emit(container, b.blocks, width_mm)
    elif isinstance(b, Box):
        _box(container, b, width_mm)
    elif isinstance(b, Table):
        _table(container, b, width_mm)
    else:  # pragma: no cover
        raise TypeError(f"unknown block {type(b)!r}")


def _box(container, b: Box, width_mm: float) -> None:
    t = container.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    _shade(cell, BOX_FILL.get(b.kind, "EAF3FB"))
    _borders(cell, BOX_EDGE.get(b.kind, "1B6CA8"), sz=10)
    if b.title:
        p = _first_par(cell)
        _runs(p, b.title, size=11, bold=True, color=RGBColor.from_string(BOX_EDGE.get(b.kind, "1B6CA8")))
    _emit(cell, b.blocks, width_mm - 6)
    if b.kind == "draw":
        for _ in range(6):
            cell.add_paragraph()
    container.add_paragraph().paragraph_format.space_after = Pt(2)


def _table(container, b: Table, width_mm: float) -> None:
    ncols = max(len(r) for r in b.rows)
    t = container.add_table(rows=len(b.rows), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    fr = b.widths or [1 / ncols] * ncols
    total = sum(fr)
    if b.width_mm:
        width_mm = min(width_mm, b.width_mm)
    widths = [Mm(width_mm * f / total) for f in fr]
    size = b.size or 10.5
    for ri, row in enumerate(b.rows):
        tr = t.rows[ri]
        if b.row_height:
            tr.height = Mm(b.row_height)
        for ci in range(ncols):
            cell = tr.cells[ci]
            cell.width = widths[ci]
            text = row[ci] if ci < len(row) else ""
            p = cell.paragraphs[0]
            head = b.header and ri == 0
            _runs(p, text, size=size, bold=head, color=RGBColor(255, 255, 255) if head else None)
            if b.align == "center" and not head:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            _borders(cell, "9FB3C8", 4)
            if head:
                _shade(cell, "1B6CA8")
            elif b.style == "stage" and ri % 2 == 0:
                _shade(cell, "F6F6F6")
            if b.bg and ri < len(b.bg) and ci < len(b.bg[ri]) and b.bg[ri][ci]:
                _shade(cell, b.bg[ri][ci].lstrip("#"))
    container.add_paragraph().paragraph_format.space_after = Pt(2)


def render_docx(doc: Doc, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(16)
    sec.top_margin, sec.bottom_margin = Mm(16), Mm(16)
    normal = d.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    width_mm = 210 - 32

    # title block
    p = d.add_paragraph()
    _runs(p, doc.title, size=20, bold=True, color=PRIMARY)
    p.paragraph_format.space_after = Pt(0)
    if doc.subtitle or doc.badge:
        p = d.add_paragraph()
        _runs(p, " · ".join(x for x in (doc.subtitle, doc.badge) if x), size=10, color=ACCENT, bold=True)
        p.paragraph_format.space_after = Pt(2)
    if doc.key:
        p = d.add_paragraph()
        _runs(p, "TEACHER'S KEY — do not print for pupils", size=10, bold=True, color=KEYRED)
    _bottom_border(d.add_paragraph())
    if doc.name_line:
        p = d.add_paragraph()
        _runs(p, "Name: ______________________________   Class: ________   Date: ______________", size=11)
        p.paragraph_format.space_after = Pt(6)

    _emit(d, doc.blocks, width_mm)

    foot = sec.footer.paragraphs[0]
    _runs(foot, doc.footer or "3-sinf · original teaching material · emoji: Twemoji (CC-BY 4.0)", size=8)
    d.core_properties.title = doc.title
    d.core_properties.author = "3-sinf"
    d.save(str(path))
    return path
